import pandas as pd
from openpyxl import load_workbook

# Load the Excel file
input_file = "./data/Material List VFORM (1).xlsx"
output_file = "./data/Processed_Material_List.xlsx"

# Open the workbook
wb = load_workbook(input_file, data_only=True)

# Dictionary to store processed data
processed_sheets = {}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    data = []
    
    category = subcategory = sub_subcategory = None
    
    for row in ws.iter_rows(values_only=False):
        first_cell = row[0]
        row_values = [cell.value for cell in row]
        
        if first_cell.value:  # It's an item or a category
            font = first_cell.font
            
            if font.size == 18 and font.color and font.color.rgb == "FF00FF00":  # Green (Category)
                category = first_cell.value
                subcategory = sub_subcategory = None
            elif font.size == 12 and font.color and font.color.rgb == "FFFFFF00":  # Yellow (Subcategory)
                subcategory = first_cell.value
                sub_subcategory = None
            elif font.size == 12 and font.color and font.color.rgb == "FF0000FF":  # Blue (Sub-Subcategory)
                sub_subcategory = first_cell.value
            else:
                # Item row: Append category, subcategory, and sub-subcategory
                row_values.extend([category, subcategory, sub_subcategory])
                data.append(row_values)

    # Convert to DataFrame and store
    df = pd.DataFrame(data)
    processed_sheets[sheet_name] = df

# Save to a new Excel file
with pd.ExcelWriter(output_file, engine="xlsxwriter") as writer:
    for sheet, df in processed_sheets.items():
        df.to_excel(writer, sheet_name=sheet, index=False)

print(f"Processed data saved to {output_file}")
